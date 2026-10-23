"""Documentos descargables: PDF (Jinja2 + WeasyPrint) y XLSX (openpyxl).

Todo documento lleva la ficha de procedencia (`ReportMeta`): fecha UTC, commit del código,
pre-registro, versión de los límites, origen de los datos, marca DEMO si aplica y el aviso
"No es asesoría financiera, legal ni fiscal". Los PDF comparten la identidad visual del
Centro de Mando (azul petróleo y acento turquesa) en versión clara para imprimir.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from typing import Any, Final

from jinja2 import Environment, PackageLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

DISCLAIMER: Final = "No es asesoría financiera, legal ni fiscal."
_ENV: Final = Environment(
    loader=PackageLoader("trading.reporting", "templates"),
    autoescape=select_autoescape(["html"]),
    trim_blocks=True,
    lstrip_blocks=True,
)


@dataclass(frozen=True, slots=True)
class ReportMeta:
    title: str
    generated_utc: str
    code_commit: str
    preregistration: str
    limits_version: str
    data_origin: str
    demo: bool
    disclaimer: str = DISCLAIMER

    @classmethod
    def create(
        cls,
        title: str,
        *,
        code_commit: str,
        preregistration: str,
        limits_version: str,
        data_origin: str,
        demo: bool,
        now: datetime | None = None,
    ) -> ReportMeta:
        moment = (now or datetime.now(UTC)).astimezone(UTC)
        return cls(
            title=title,
            generated_utc=moment.strftime("%Y-%m-%d %H:%M UTC"),
            code_commit=code_commit,
            preregistration=preregistration,
            limits_version=limits_version,
            data_origin=data_origin,
            demo=demo,
        )

    def lines(self) -> list[tuple[str, str]]:
        return [
            ("Generado", self.generated_utc),
            ("Commit del código", self.code_commit),
            ("Pre-registro", self.preregistration),
            ("Versión de los límites", self.limits_version),
            ("Origen de los datos", self.data_origin),
            ("Datos de demostración", "SÍ" if self.demo else "no"),
            ("Aviso", self.disclaimer),
        ]


def render_html(template: str, meta: ReportMeta, context: Mapping[str, Any]) -> str:
    return _ENV.get_template(template).render(meta=meta, meta_lines=meta.lines(), **context)


def render_pdf(template: str, meta: ReportMeta, context: Mapping[str, Any]) -> bytes:
    from weasyprint import HTML  # importación diferida: carga Pango solo al generar un PDF

    pdf: bytes = HTML(string=render_html(template, meta, context)).write_pdf()
    return pdf


def svg_line(values: Sequence[float], *, width: int = 640, height: int = 160) -> str:
    """Polilínea SVG simple para curvas en los PDF (sin librerías de gráficos)."""
    if len(values) < 2:
        return ""
    low, high = min(values), max(values)
    span = high - low or 1.0
    step = width / (len(values) - 1)
    points = " ".join(
        f"{i * step:.1f},{height - (v - low) / span * (height - 8) - 4:.1f}"
        for i, v in enumerate(values)
    )
    return (
        f'<svg viewBox="0 0 {width} {height}" width="100%" height="{height}" '
        f'xmlns="http://www.w3.org/2000/svg" role="img">'
        f'<polyline fill="none" stroke="#0f9d8a" stroke-width="2" points="{points}"/></svg>'
    )


def workbook(
    meta: ReportMeta, sheets: Mapping[str, tuple[Sequence[str], Sequence[Sequence[Any]]]]
) -> bytes:
    """XLSX con una hoja de procedencia y una hoja por tabla."""
    book = Workbook()
    info = book.active
    assert info is not None  # noqa: S101 - un libro nuevo siempre tiene una hoja activa
    info.title = "Procedencia"
    info.append([meta.title])
    info["A1"].font = Font(bold=True, size=14)
    if meta.demo:
        info.append(["DATOS DE DEMOSTRACIÓN: no son operaciones reales"])
        info["A2"].font = Font(bold=True, color="B45309")
    for label, value in meta.lines():
        info.append([label, value])
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 80
    header_fill = PatternFill("solid", fgColor="0B2A33")
    for name, (header, rows) in sheets.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(list(header))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(list(row))
        for index, column in enumerate(header, start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            sheet.column_dimensions[letter].width = max(12, min(40, len(str(column)) + 4))
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def meta_dict(meta: ReportMeta) -> dict[str, Any]:
    return asdict(meta)
