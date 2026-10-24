"""Centro de Mando: rutas, honestidad de los datos y garantía de solo lectura."""

import io
import json
import zipfile
from pathlib import Path

import polars as pl
import pytest
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api_fixtures import build_repo, client_for
from trading.api.app import WRITE_ROUTES, create_app
from trading.api.settings import ApiSettings

READ_ONLY = {"GET", "HEAD"}


def test_the_api_has_no_write_route_on_money_risk_or_execution(repo: Path) -> None:
    app = create_app(
        ApiSettings(
            auth_enabled=False, repo_root=repo, data_root=repo / "data", app_state=repo / "app"
        )
    )
    writes = {
        (method, route.path)
        for route in app.routes
        if isinstance(route, APIRoute)
        for method in (route.methods or set()) - READ_ONLY
    }
    assert writes == WRITE_ROUTES
    forbidden = ("orden", "order", "live", "clave", "key", "ejecucion", "retiro", "withdraw")
    for method, path in writes:
        assert not any(word in path.lower() for word in forbidden), (method, path)
    assert all(not path.startswith("/api/riesgo") for _, path in writes)


def test_status_today_and_brief(client: TestClient) -> None:
    status = client.get("/api/estado").json()
    assert status["live_trading"] is False
    assert status["holdout"]["sellado"] is True
    assert status["intentos"] == {"principales": 2, "total": 4}
    todo = status["pendientes_propietario"]
    assert todo[0]["documento"] == "docs/f1-verificacion-plataforma.md"
    assert all(t["tarea"] and t["documento"] for t in todo)
    today = client.get("/api/hoy").json()
    assert today["cuenta"]["origen"] == "demo"
    assert today["parte_del_dia"]["validado"] is True
    assert "LIVE_TRADING sigue apagado" in today["parte_del_dia"]["texto"]
    assert today["posiciones"]["items"] == []


def test_laboratory(client: TestClient) -> None:
    lab = client.get("/api/laboratorio").json()
    ids = [h["id"] for h in lab["hipotesis"]]
    assert ids == ["H1-v1", "H2-v1", "H4-H1-v1", "H4-H2-v1", "H3"]
    h1 = lab["hipotesis"][0]
    assert h1["ciclo"][3] == {"etapa": "gate F4", "estado": "rechazado"}
    assert "le faltan 194 operaciones" in h1["pista"]
    detail = client.get("/api/laboratorio/H1-v1").json()
    assert detail["origen"] == "backtest"
    assert len(detail["curva"]) == 30
    assert min(p["dd"] for p in detail["underwater"]) < 0
    assert sum(b["n"] for b in detail["distribucion_r"]) == 6
    assert detail["abanico_monte_carlo"][0]["p5"] <= detail["abanico_monte_carlo"][0]["p95"]
    lookback = next(s for s in detail["sensibilidad"] if s["parametro"] == "lookback")
    assert [v["valor"] for v in lookback["valores"]] == [6]
    assert {v["codigo"] for v in detail["vetos"]} == {"RISK_PER_TRADE", "BELOW_MIN_NOTIONAL"}
    assert client.get("/api/laboratorio/H4-H1-v1").status_code == 404
    assert client.get("/api/laboratorio/H9").status_code == 404


def test_market(client: TestClient) -> None:
    assert client.get("/api/mercado/simbolos").json()["simbolos"] == ["BTCUSDT", "ETHUSDT"]
    four = client.get(
        "/api/mercado/velas", params={"simbolo": "BTCUSDT", "intervalo": "4h", "dias": 30}
    ).json()
    assert four["origen"] == "real"
    assert len(four["velas"]) == 180
    assert set(four["indicadores"]) >= {"canal_alto", "atr100"}
    long_range = {"simbolo": "BTCUSDT", "intervalo": "4h", "dias": 400}
    assert client.get("/api/mercado/velas", params=long_range).json()["marcadores"]
    assert four["funding"][0]["tasa_8h"] == pytest.approx(0.0001)
    daily = client.get(
        "/api/mercado/velas", params={"simbolo": "BTCUSDT", "intervalo": "1d", "dias": 60}
    ).json()
    assert "senal_h1" in daily["indicadores"]
    missing = client.get("/api/mercado/velas", params={"simbolo": "SOLUSDT", "intervalo": "1d"})
    assert missing.status_code == 404
    assert (
        client.get("/api/mercado/velas", params={"simbolo": "x", "intervalo": "1d"}).status_code
        == 422
    )
    assert client.get("/api/mercado/regimenes").json()["tramos"]
    corr = client.get("/api/mercado/correlaciones").json()
    assert corr["matriz"][0][0] == pytest.approx(1.0)


def test_market_without_data_is_demo(tmp_path: Path) -> None:
    client = client_for(build_repo(tmp_path, with_data=False))
    assert client.get("/api/mercado/simbolos").json()["origen"] == "demo"
    candles = client.get(
        "/api/mercado/velas", params={"simbolo": "BTCUSDT", "intervalo": "4h"}
    ).json()
    assert candles["origen"] == "demo"
    assert client.get("/api/mercado/regimenes").json()["origen"] == "demo"
    assert client.get("/api/mercado/correlaciones").json()["origen"] == "demo"
    assert client.get("/api/estado").json()["datos"]["disponible"] is False


def test_risk_calculator_uses_the_real_engine(client: TestClient) -> None:
    ok = client.get(
        "/api/riesgo/calculadora",
        params={"simbolo": "BTCUSDT", "lado": "LONG", "entrada": "60000", "stop": "58800"},
    ).json()
    assert ok["aprobada"] is True
    assert ok["vetos"] == []
    assert ok["presupuesto_usd"] == "7.500"
    assert float(ok["riesgo_usd"]) <= 7.5
    vetoed = client.get(
        "/api/riesgo/calculadora",
        params={
            "simbolo": "SOLUSDT",
            "lado": "SHORT",
            "entrada": "150",
            "stop": "160",
            "mercado": "SPOT",
        },
    ).json()
    codes = {v["code"] for v in vetoed["vetos"]}
    assert {"SHORT_ON_SPOT", "SHORT_NOT_VALIDATED"} <= codes
    assert all(v["message"] for v in vetoed["vetos"])
    for bad in ({"entrada": "abc", "stop": "1"}, {"entrada": "-1", "stop": "1"}):
        assert client.get("/api/riesgo/calculadora", params=bad).status_code == 422
    limits = client.get("/api/riesgo/limites").json()
    assert limits["editable"] is False
    assert limits["vigentes"][0] == {
        "nombre": "Riesgo por operación",
        "valor": "0.005",
        "techo_plan": "0.01",
        "unidad": "fracción",
        "adr": "ADR-007",
    }
    assert limits["propuestos_adr_013"]["macro_blackout"] == "120 min"
    history = client.get("/api/riesgo/vetos").json()
    assert history["vetos"][0] == {
        "codigo": "RISK_PER_TRADE",
        "titulo": "Riesgo por operación",
        "explicacion": "Esta operación arriesgaría más que tu límite por operación.",
        "veces": 4,
    }


def test_decisions_project_glossary(client: TestClient) -> None:
    decisions = client.get("/api/decisiones").json()
    assert decisions["origen"] == "demo"
    assert len(decisions["libros"]) == 3
    project = client.get("/api/proyecto").json()
    assert project["fases"][0]["id"] == "F0"
    assert any(a["id"] == "013" and "Propuesto" in a["estado"] for a in project["adr"])
    assert any(r["id"] == "R01" for r in project["riesgos"])
    assert project["costes"]["conceptos"][0]["concepto"] == "Claude Max 5x"
    assert (
        project["pendientes_propietario"]
        == client.get("/api/estado").json()["pendientes_propietario"]
    )
    glossary = client.get("/api/glosario").json()
    assert glossary["copiloto"] == "Brújula"
    assert "PBO" in glossary["terminos"]


def test_journal_and_notes(client: TestClient) -> None:
    journal = client.get("/api/diario").json()
    assert journal["origen"] == "backtest"
    assert len(journal["operaciones"]) == 12
    assert client.post("/api/diario/notas", json={"texto": "sin cabecera"}).status_code == 403
    headers = {"x-brujula": "1"}
    created = client.post(
        "/api/diario/notas",
        json={"referencia": "H1-v1-0001", "texto": " revisar stop "},
        headers=headers,
    )
    assert created.status_code == 200
    assert created.json()["texto"] == "revisar stop"
    assert client.post("/api/diario/notas", json={"texto": ""}, headers=headers).status_code == 422
    assert (
        client.post("/api/diario/notas", json={"texto": "   "}, headers=headers).status_code == 422
    )
    assert [n["referencia"] for n in client.get("/api/diario").json()["notas"]] == ["H1-v1-0001"]


def test_security_headers(client: TestClient) -> None:
    response = client.get("/api/estado")
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["referrer-policy"] == "no-referrer"


# --- reportes ------------------------------------------------------------------------------


def _get(client: TestClient, kind: str, fmt: str, **params: str) -> bytes:
    response = client.get(f"/api/reportes/{kind}/{fmt}", params=params)
    assert response.status_code == 200, response.text
    return bytes(response.content)


def test_pdf_reports_open(client: TestClient) -> None:
    for kind in ("semanal", "mensual"):
        content = _get(client, kind, "pdf")
        assert content.startswith(b"%PDF")
    tear = _get(client, "ficha", "pdf", hipotesis="H1-v1")
    assert tear.startswith(b"%PDF")
    assert (
        client.get("/api/reportes/ficha/pdf", params={"hipotesis": "H4-H1-v1"}).status_code == 404
    )
    assert client.get("/api/reportes/semanal/docx").status_code == 404
    history = client.get("/api/reportes").json()["historial"]
    assert [h["tipo"] for h in history][:3] == ["ficha", "mensual", "semanal"]
    assert history[1]["demo"] is True


def test_accountant_totals_match_the_trades(client: TestClient) -> None:
    trades = client.get("/api/diario").json()["operaciones"]
    expected = sum(float(t["pnl"]) for t in trades)
    book = load_workbook(io.BytesIO(_get(client, "contador", "xlsx")))
    assert book.sheetnames == ["Procedencia", "Operaciones", "Totales"]
    assert book["Procedencia"]["A2"].value.startswith("DATOS DE DEMOSTRACIÓN")
    totals = {row[0]: row[1] for row in book["Totales"].iter_rows(min_row=2, values_only=True)}
    assert float(totals["resultado_neto_usd"]) == pytest.approx(expected, abs=0.01 * len(trades))
    assert totals["resultado_neto_mxn"] == "por completar"
    rows = list(book["Operaciones"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(trades)
    text = _get(client, "contador", "csv").decode("utf-8")
    assert text.startswith("DATOS DE DEMOSTRACIÓN")
    assert text.count("por completar") == 3 * len(trades)


def test_monthly_workbook_and_data_exports(client: TestClient) -> None:
    book = load_workbook(io.BytesIO(_get(client, "mensual", "xlsx")))
    assert "Hipótesis" in book.sheetnames
    csv_text = _get(client, "datos", "csv").decode("utf-8")
    assert csv_text.splitlines()[0].startswith("id,hipotesis")
    frame = pl.read_parquet(io.BytesIO(_get(client, "datos", "parquet")))
    assert frame.height == 12


def test_audit_export_has_a_valid_chain(client: TestClient) -> None:
    payload = json.loads(_get(client, "auditoria", "json"))
    assert payload["verificacion"] == {"valida": True, "entradas": 18}
    rows = _get(client, "auditoria", "csv").decode("utf-8").splitlines()
    assert rows[0] == "indice,fecha_utc,tipo,datos,hash_anterior,hash"
    assert len(rows) == 19


def test_spa_is_served_with_fallback(tmp_path: Path) -> None:
    repo = build_repo(tmp_path, with_data=False)
    dist = repo / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html>app</html>", encoding="utf-8")
    (dist / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    (dist / "favicon.svg").write_text("<svg/>", encoding="utf-8")
    client = client_for(repo)
    assert client.get("/laboratorio").text == "<html>app</html>"
    assert client.get("/favicon.svg").text == "<svg/>"
    assert client.get("/assets/app.js").text == "console.log(1)"
    assert "frame-ancestors" in client.get("/").headers["content-security-policy"]


def test_zip_magic_of_xlsx(client: TestClient) -> None:
    assert zipfile.is_zipfile(io.BytesIO(_get(client, "mensual", "xlsx")))
